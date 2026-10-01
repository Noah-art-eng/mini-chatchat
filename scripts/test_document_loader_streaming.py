"""验证文档路径到路径解析及生产上传流程对该接口的接入。"""

import atexit
import inspect
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
from docx import Document


ROOT_DIR = Path(__file__).resolve().parents[1]
BACKEND_DIR = ROOT_DIR / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

TEST_RUNTIME = tempfile.TemporaryDirectory()
atexit.register(TEST_RUNTIME.cleanup)
os.environ["MINI_CHATCHAT_DB_PATH"] = str(Path(TEST_RUNTIME.name) / "mini.db")
os.environ["MINI_CHATCHAT_DATA_ROOT"] = str(Path(TEST_RUNTIME.name) / "data")
os.environ["MINI_CHATCHAT_UPLOADS_DIR"] = str(Path(TEST_RUNTIME.name) / "uploads")

from services import document_loader  # noqa: E402
import rag  # noqa: E402


class FakeEmbeddingModel:
    """让生产接入测试不依赖真实 Embedding 模型下载与推理。"""

    def encode(self, texts):
        """返回足够初始化隔离测试服务的固定向量。"""
        return np.ones((len(texts), 2), dtype="float32")


import services.kb_service as kb_service_module  # noqa: E402

kb_service_module.get_embedding_model = lambda _name: FakeEmbeddingModel()

import app as backend_app  # noqa: E402
import chat_service  # noqa: E402
from auth.models import guest_user  # noqa: E402


class RecordingReader:
    """记录解析器每次读取的大小，用于防止新路径退回无参数完整读取。"""

    def __init__(self, file_object):
        """保存真实文本流，同时收集后续 read() 收到的 size。"""
        self.file_object = file_object
        self.read_sizes = []

    def read(self, size=-1):
        """记录读取大小后交给真实文本流处理。"""
        self.read_sizes.append(size)
        return self.file_object.read(size)

    def __enter__(self):
        """让记录包装器可以沿用解析器的 with 打开方式。"""
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        """退出读取范围时关闭底层真实文件。"""
        self.file_object.close()


class FailingReader(RecordingReader):
    """在成功返回第一块后抛出指定异常，模拟解析中途读取失败。"""

    def __init__(self, file_object, error):
        """保存要原样抛出的异常，并初始化读取次数。"""
        super().__init__(file_object)
        self.error = error
        self.read_count = 0

    def read(self, size=-1):
        """第一次读取真实内容，第二次开始抛出测试指定的异常。"""
        self.read_count += 1
        if self.read_count > 1:
            raise self.error
        return super().read(size)


class DocumentLoaderStreamingTest(unittest.TestCase):
    """验证五种受支持文档的路径到路径解析契约。"""

    def test_txt_content_is_preserved_exactly(self):
        """确认普通 TXT 经过路径到路径解析后内容与旧读取语义一致。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            destination = Path(temp_dir) / "content.txt"
            content = "first line\nsecond line\n"
            source.write_text(content, encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), content)

    def test_markdown_content_is_preserved_exactly(self):
        """确认 Markdown 标记和空白不会在增量复制时被改写。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.md"
            destination = Path(temp_dir) / "content.txt"
            content = "# Title\n\n- one\n- two  \n"
            source.write_text(content, encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), content)

    def test_small_chunks_preserve_chinese_english_and_emoji(self):
        """确认很小的文本块也不会切坏中文、英文或 emoji 字符。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "混合内容.txt"
            destination = Path(temp_dir) / "解析结果.txt"
            content = "中文 Mini ChatChat 🙂\n第二行🚀end"
            source.write_text(content, encoding="utf-8")

            document_loader.parse_file_to_text_file(
                source,
                destination,
                chunk_chars=2,
            )

            self.assertEqual(destination.read_text(encoding="utf-8"), content)

    def test_parser_uses_multiple_sized_reads(self):
        """确认较长文本通过多次 read(size) 处理，且从不调用无参数 read()。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            destination = Path(temp_dir) / "content.txt"
            content = "abcdefghij"
            source.write_text(content, encoding="utf-8")
            real_open = open
            recording_reader = RecordingReader(
                real_open(source, "r", encoding="utf-8")
            )

            def open_for_test(file_path, *args, **kwargs):
                """只包装 source 的读取，staging 写入仍使用真实文件系统。"""
                if os.fspath(file_path) == os.fspath(source):
                    return recording_reader
                return real_open(file_path, *args, **kwargs)

            with patch.object(
                document_loader,
                "open",
                side_effect=open_for_test,
                create=True,
            ):
                document_loader.parse_file_to_text_file(
                    source,
                    destination,
                    chunk_chars=3,
                )

            self.assertEqual(destination.read_text(encoding="utf-8"), content)
            self.assertGreater(len(recording_reader.read_sizes), 2)
            self.assertEqual(set(recording_reader.read_sizes), {3})

    def test_empty_file_creates_empty_destination(self):
        """确认空 TXT 仍生成空结果文件，不擅自改变旧有空文件行为。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "empty.txt"
            destination = Path(temp_dir) / "content.txt"
            source.write_text("", encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), "")

    def test_existing_destination_is_replaced_after_success(self):
        """确认完整解析成功后才用新内容替换已有 destination。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            destination = Path(temp_dir) / "content.txt"
            source.write_text("new content", encoding="utf-8")
            destination.write_text("old content", encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(
                destination.read_text(encoding="utf-8"),
                "new content",
            )
            self.assertEqual(
                sorted(path.name for path in Path(temp_dir).iterdir()),
                ["content.txt", "source.txt"],
            )

    def test_missing_destination_directory_is_created(self):
        """确认新 API 会先建立 content 目标目录，再在该目录内完成 staging。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            destination = Path(temp_dir) / "kb" / "content" / "source.txt"
            source.write_text("content", encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), "content")
            self.assertEqual(
                [path.name for path in destination.parent.iterdir()],
                ["source.txt"],
            )

    def test_read_failure_preserves_destination_and_cleans_staging(self):
        """确认解析中途失败会保留旧结果、清理 staging，并原样抛出异常。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            destination = Path(temp_dir) / "content.txt"
            source.write_text("new content", encoding="utf-8")
            destination.write_text("old content", encoding="utf-8")
            real_open = open
            expected_error = RuntimeError("injected read failure")

            def open_for_test(file_path, *args, **kwargs):
                """让 source 第二次读取失败，其余文件操作保持真实。"""
                if os.fspath(file_path) == os.fspath(source):
                    return FailingReader(
                        real_open(source, "r", encoding="utf-8"),
                        expected_error,
                    )
                return real_open(file_path, *args, **kwargs)

            with (
                patch.object(
                    document_loader,
                    "open",
                    side_effect=open_for_test,
                    create=True,
                ),
                self.assertRaises(RuntimeError) as raised,
            ):
                document_loader.parse_file_to_text_file(
                    source,
                    destination,
                    chunk_chars=3,
                )

            self.assertIs(raised.exception, expected_error)
            self.assertEqual(destination.read_text(encoding="utf-8"), "old content")
            self.assertEqual(
                sorted(path.name for path in Path(temp_dir).iterdir()),
                ["content.txt", "source.txt"],
            )

    def test_replace_failure_preserves_destination_and_cleans_staging(self):
        """确认最终切换失败时旧结果不变，且完整但未安装的 staging 被删除。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            destination = Path(temp_dir) / "content.txt"
            source.write_text("new content", encoding="utf-8")
            destination.write_text("old content", encoding="utf-8")
            expected_error = OSError("injected replace failure")

            with (
                patch.object(
                    document_loader.os,
                    "replace",
                    side_effect=expected_error,
                ),
                self.assertRaises(OSError) as raised,
            ):
                document_loader.parse_file_to_text_file(source, destination)

            self.assertIs(raised.exception, expected_error)
            self.assertEqual(destination.read_text(encoding="utf-8"), "old content")
            self.assertEqual(
                sorted(path.name for path in Path(temp_dir).iterdir()),
                ["content.txt", "source.txt"],
            )

    def test_csv_rows_preserve_legacy_output(self):
        """确认 CSV 逐行写入后仍使用竖线连接字段，并且末尾不增加换行。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.csv"
            destination = Path(temp_dir) / "content.txt"
            source.write_text("name,score\nAlice,10\nBob,20\n", encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(
                destination.read_text(encoding="utf-8"),
                "name | score\nAlice | 10\nBob | 20",
            )

    def test_empty_csv_creates_empty_destination(self):
        """确认空 CSV 与旧 loader 一样输出空文本。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "empty.csv"
            destination = Path(temp_dir) / "content.txt"
            source.write_text("", encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), "")

    def test_csv_uses_standard_quoted_field_parsing(self):
        """确认带逗号的 quoted field 仍由 csv.reader 按旧规则解析。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "quoted.csv"
            destination = Path(temp_dir) / "content.txt"
            source.write_text('name,note\nAlice,"hello, world"\n', encoding="utf-8")

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(
                destination.read_text(encoding="utf-8"),
                "name | note\nAlice | hello, world",
            )

    def test_pdf_pages_preserve_load_pdf_output(self):
        """确认 PDF 逐页写入仍严格保留 rag.load_pdf() 的页尾换行规则。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.pdf"
            destination = Path(temp_dir) / "content.txt"
            source.write_bytes(b"placeholder")

            class Page:
                """提供指定的 extract_text() 结果。"""

                def __init__(self, text):
                    self.text = text

                def extract_text(self):
                    return self.text

            pages = [Page("page one"), Page(None), Page("page two\n")]
            with patch.object(rag, "PdfReader", return_value=SimpleNamespace(pages=pages)):
                legacy_output = rag.load_pdf(source)

            pages = [Page("page one"), Page(None), Page("page two\n")]
            with patch.object(
                document_loader,
                "PdfReader",
                return_value=SimpleNamespace(pages=pages),
            ):
                document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), legacy_output)

    def test_pdf_pages_without_text_create_empty_destination(self):
        """确认 PDF 空页和 extract_text() 返回 None 时不会写入占位换行。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "empty.pdf"
            destination = Path(temp_dir) / "content.txt"
            source.write_bytes(b"placeholder")
            pages = [
                SimpleNamespace(extract_text=lambda: None),
                SimpleNamespace(extract_text=lambda: ""),
            ]

            with patch.object(
                document_loader,
                "PdfReader",
                return_value=SimpleNamespace(pages=pages),
            ):
                document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), "")

    def test_docx_paragraphs_preserve_legacy_output(self):
        """确认 DOCX 段落逐个写入后仍以换行连接，末尾不增加换行。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.docx"
            destination = Path(temp_dir) / "content.txt"
            document = Document()
            document.add_paragraph("first")
            document.add_paragraph("second")
            document.save(source)

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(
                destination.read_text(encoding="utf-8"),
                "first\nsecond",
            )

    def test_docx_filters_empty_paragraphs_like_legacy_loader(self):
        """确认 DOCX 空段落继续被过滤，不改变旧 load_docx() 语义。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.docx"
            destination = Path(temp_dir) / "content.txt"
            document = Document()
            document.add_paragraph("")
            document.add_paragraph("kept")
            document.add_paragraph("")
            document.save(source)

            document_loader.parse_file_to_text_file(source, destination)

            self.assertEqual(destination.read_text(encoding="utf-8"), "kept")

    def test_structured_parser_failures_preserve_destination_and_cleanup(self):
        """确认 CSV/PDF/DOCX 中途失败均保留旧结果、清理 staging 并抛原异常。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)

            for extension in (".csv", ".pdf", ".docx"):
                with self.subTest(extension=extension):
                    source = root / f"source{extension}"
                    destination = root / f"content-{extension[1:]}.txt"
                    source.write_bytes(b"placeholder")
                    destination.write_text("old content", encoding="utf-8")
                    expected_error = RuntimeError(f"{extension} parse failed")

                    if extension == ".csv":
                        def failing_rows():
                            yield ["partial"]
                            raise expected_error

                        format_patch = patch.object(
                            document_loader.csv,
                            "reader",
                            return_value=failing_rows(),
                        )
                    elif extension == ".pdf":
                        pages = [
                            SimpleNamespace(extract_text=lambda: "partial"),
                            SimpleNamespace(
                                extract_text=lambda error=expected_error: (_ for _ in ()).throw(error)
                            ),
                        ]
                        format_patch = patch.object(
                            document_loader,
                            "PdfReader",
                            return_value=SimpleNamespace(pages=pages),
                        )
                    else:
                        class FailingParagraphs:
                            def __iter__(self):
                                yield SimpleNamespace(text="partial")
                                raise expected_error

                        format_patch = patch.object(
                            document_loader,
                            "Document",
                            return_value=SimpleNamespace(paragraphs=FailingParagraphs()),
                        )

                    before = sorted(path.name for path in root.iterdir())
                    with format_patch, self.assertRaises(RuntimeError) as raised:
                        document_loader.parse_file_to_text_file(source, destination)

                    self.assertIs(raised.exception, expected_error)
                    self.assertEqual(
                        destination.read_text(encoding="utf-8"),
                        "old content",
                    )
                    self.assertEqual(
                        sorted(path.name for path in root.iterdir()),
                        before,
                    )

    def test_cleanup_failure_does_not_replace_original_parse_error(self):
        """确认 staging 删除失败不会掩盖更早发生的文档解析错误。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.pdf"
            destination = Path(temp_dir) / "content.txt"
            source.write_bytes(b"placeholder")
            expected_error = RuntimeError("parse failed")
            page = SimpleNamespace(
                extract_text=lambda: (_ for _ in ()).throw(expected_error)
            )

            with (
                patch.object(
                    document_loader,
                    "PdfReader",
                    return_value=SimpleNamespace(pages=[page]),
                ),
                patch.object(document_loader.os, "remove", side_effect=OSError("cleanup failed")),
                self.assertRaises(RuntimeError) as raised,
            ):
                document_loader.parse_file_to_text_file(source, destination)

            self.assertIs(raised.exception, expected_error)

    def test_destination_without_directory_uses_current_directory(self):
        """确认只有文件名的 destination 会在当前目录安全完成 staging 和替换。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.txt"
            source.write_text("content", encoding="utf-8")
            previous_directory = os.getcwd()
            try:
                os.chdir(temp_dir)
                document_loader.parse_file_to_text_file(
                    source,
                    "content.txt",
                    chunk_chars=2,
                )
            finally:
                os.chdir(previous_directory)

            self.assertEqual(
                (Path(temp_dir) / "content.txt").read_text(encoding="utf-8"),
                "content",
            )


class ProductionParserIntegrationTest(unittest.TestCase):
    """验证普通上传、reindex 和 temp KB 已切换到路径到路径解析。"""

    def test_normal_upload_worker_uses_path_to_path_parser(self):
        """普通上传的同步 worker 应直接生成 content 文件，不再接收完整正文。"""
        service = SimpleNamespace(
            chunk_size=None,
            chunk_overlap=None,
            mutation_lock=threading.RLock(),
            rebuild_and_sync=unittest.mock.Mock(),
        )

        with patch.object(backend_app, "parse_file_to_text_file") as parser:
            backend_app.process_uploaded_document(
                service,
                "/tmp/raw.pdf",
                "/tmp/content.txt",
                300,
                50,
            )

        parser.assert_called_once_with("/tmp/raw.pdf", "/tmp/content.txt")
        service.rebuild_and_sync.assert_called_once()

    def test_reindex_uses_path_to_path_parser(self):
        """文档 reindex 应重新解析 raw upload，而不是先构造完整 Python 字符串。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "uploads" / "report.pdf"
            content_dir = root / "content"
            source.parent.mkdir()
            content_dir.mkdir()
            source.write_bytes(b"placeholder")
            service = SimpleNamespace(
                kb_name="default",
                content_path=str(content_dir),
                chunk_size=None,
                chunk_overlap=None,
                mutation_lock=threading.RLock(),
                rebuild_and_sync=unittest.mock.Mock(),
            )

            def parse_for_test(source_path, destination_path):
                Path(destination_path).write_text("parsed", encoding="utf-8")

            with (
                patch.object(backend_app, "require_permission"),
                patch.object(backend_app, "get_scoped_kb_service", return_value=service),
                patch.object(backend_app, "find_reindex_source_file", return_value=str(source)),
                patch.object(backend_app, "update_file_status"),
                patch.object(backend_app, "delete_file_docs"),
                patch.object(
                    backend_app,
                    "parse_file_to_text_file",
                    side_effect=parse_for_test,
                ) as parser,
            ):
                result = backend_app.reindex_document(
                    "report.pdf",
                    backend_app.ReindexFileRequest(),
                    guest_user(),
                )

            self.assertEqual(result["filename"], "report.txt")
            parser.assert_called_once_with(str(source), str(content_dir / "report.txt"))
            service.rebuild_and_sync.assert_called_once()

    def test_temp_kb_uses_path_to_path_parser(self):
        """临时文件问答应从已安装 raw 文件直接生成 content 文本文件。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            staged = root / "staged.docx"
            staged.write_bytes(b"placeholder")
            temp_root = root / "temp-kbs"
            fake_service = object()

            def parse_for_test(source_path, destination_path):
                Path(destination_path).write_text("parsed", encoding="utf-8")

            with (
                patch.object(chat_service, "get_temp_root_path", return_value=str(temp_root)),
                patch.object(chat_service, "migrate_legacy_demo_files"),
                patch.object(
                    chat_service,
                    "parse_file_to_text_file",
                    side_effect=parse_for_test,
                ) as parser,
                patch.object(chat_service, "MiniKBService", return_value=fake_service),
            ):
                result = chat_service.create_temp_kb_from_upload(
                    str(staged),
                    "资料.docx",
                    user_id=123,
                )

            temp_kb_id = result["temp_kb_id"]
            expected_root = temp_root / temp_kb_id
            parser.assert_called_once()
            actual_source, actual_destination = parser.call_args.args
            self.assertEqual(
                os.path.realpath(actual_source),
                os.path.realpath(expected_root / "uploads" / "资料.docx"),
            )
            self.assertEqual(
                os.path.realpath(actual_destination),
                os.path.realpath(expected_root / "content" / "资料.txt"),
            )
            self.assertIs(chat_service.temp_kb_services[(123, temp_kb_id)], fake_service)
            chat_service.temp_kb_services.pop((123, temp_kb_id), None)

    def test_production_functions_do_not_call_legacy_load_file(self):
        """锁定三个生产解析入口不再恢复为 load_file() 加一次性 write。"""
        functions = (
            backend_app.process_uploaded_document,
            backend_app.reindex_document,
            chat_service.create_temp_kb_from_upload,
        )
        for function in functions:
            with self.subTest(function=function.__name__):
                self.assertNotIn("load_file(", inspect.getsource(function))


if __name__ == "__main__":
    unittest.main()
